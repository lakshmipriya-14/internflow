from io import BytesIO

from docx import Document
from fastapi import APIRouter, File, HTTPException, UploadFile
from pypdf import PdfReader
from app.core.deps import CurrentUser, DbSession
from app.schemas.profile import ProfileResponse, ProfileUpdate, ResumeParseRequest, ResumeParseResponse
from app.services import profile_service, ai_service

router = APIRouter(prefix="/profile", tags=["Student Profile"])
MAX_RESUME_SIZE = 10 * 1024 * 1024


def _extract_resume_text(filename: str, content: bytes) -> str:
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    try:
        if suffix == "pdf":
            reader = PdfReader(BytesIO(content), strict=False)
            extracted_pages = []
            for page in reader.pages:
                try:
                    p_text = page.extract_text()
                    if p_text:
                        extracted_pages.append(p_text.strip())
                except Exception:
                    continue
            text = "\n\n".join(extracted_pages)
        elif suffix == "docx":
            document = Document(BytesIO(content))
            parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
            parts.extend(
                " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                for table in document.tables
                for row in table.rows
            )
            text = "\n".join(parts)
        elif suffix in {"txt", "md"}:
            for encoding in ("utf-8-sig", "utf-8", "latin-1", "cp1252"):
                try:
                    text = content.decode(encoding)
                    break
                except UnicodeDecodeError:
                    continue
            else:
                text = content.decode("utf-8", errors="replace")
        elif suffix == "doc":
            raise HTTPException(
                status_code=415,
                detail="Legacy .doc format is not supported. Please save your file as .docx or .pdf, or paste the text directly."
            )
        else:
            raise HTTPException(
                status_code=415,
                detail="Unsupported format. Upload a PDF, DOCX, TXT, or Markdown resume."
            )
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail="Could not read this resume file. Ensure it is a valid, uncorrupted document."
        ) from exc

    if not text.strip():
        raise HTTPException(
            status_code=400,
            detail="No readable text was found in this file. If this is a scanned/image PDF, please export it with selectable text or copy-paste the text directly."
        )
    if len(text) > 100_000:
        raise HTTPException(status_code=413, detail="Extracted resume text must be 100,000 characters or fewer.")
    return text.strip()


def _save_parsed_resume(db, user_id: int, resume_text: str):
    result = ai_service.parse_resume_text(resume_text)
    updated_profile = profile_service.update_profile(
        db,
        user_id,
        ProfileUpdate(
            resume_text=resume_text,
            skills=", ".join(result["extracted_skills"]) or None,
            bio=result["suggested_bio"],
            target_roles=", ".join(result["suggested_roles"]) or None
        )
    )
    return result, updated_profile


@router.get("", response_model=ProfileResponse)
def get_profile(current_user: CurrentUser, db: DbSession):
    """Retrieve current student profile."""
    return profile_service.get_profile_by_user_id(db, current_user.id)


@router.put("", response_model=ProfileResponse)
def update_profile(payload: ProfileUpdate, current_user: CurrentUser, db: DbSession):
    """Update current student profile details."""
    return profile_service.update_profile(db, current_user.id, payload)


@router.post("/parse-resume", response_model=ResumeParseResponse)
def parse_resume(payload: ResumeParseRequest, current_user: CurrentUser, db: DbSession):
    """Parse resume text and automatically extract skills and suggested bio/roles."""
    if not payload.resume_text.strip():
        raise HTTPException(status_code=400, detail="Resume text cannot be empty.")
    result, _ = _save_parsed_resume(db, current_user.id, payload.resume_text.strip())
    return result


@router.post("/upload-resume")
async def upload_resume(current_user: CurrentUser, db: DbSession, file: UploadFile = File(...)):
    """Extract, parse, and save a user's uploaded PDF, DOCX, or text resume."""
    content = await file.read(MAX_RESUME_SIZE + 1)
    if len(content) > MAX_RESUME_SIZE:
        raise HTTPException(status_code=413, detail="Resume file must be 10 MB or smaller.")
    resume_text = _extract_resume_text(file.filename or "", content)
    result, updated_profile = _save_parsed_resume(db, current_user.id, resume_text)
    return {
        **result,
        "resume_text": resume_text,
        "profile": ProfileResponse.model_validate(updated_profile),
    }
